"""
Repara el desfase entre disco y ColaRevision detectado el 2026-08-04:
~272 PDFs en corregir_manualmente/ e incidencias/ que se movieron a su
carpeta pero cuya fila en ColaRevision nunca se guardó (fallo silencioso de
guardar_en_cola_revision_sql, ver sql_historial.py), y 2 filas en
ColaRevision de pendientes cuyo PDF ya no existe en la carpeta.

Para cada huérfano en disco intenta recuperar sus datos ya extraídos del
historial (historial_facturas_auto.xlsx, donde se guardan siempre
independientemente de si el guardado en SQL falló); si no aparece ahí
(p.ej. incidencias por "varias facturas en un PDF", que nunca llegan a
generar una fila con datos), inserta un placeholder -mismo criterio que ya
usa mover_pdf en logic.py cuando no hay fila- para que al menos aparezca en
su pestaña y se pueda completar a mano.

Escribe en SQL Server de producción (INSERT/UPDATE en ColaRevision,
DELETE de las 2 filas fantasma). Pensado para ejecutarse una sola vez.
"""

import os

from openpyxl import load_workbook

from logic import CORREGIR_DIR, INCIDENCIAS_DIR, HISTORIAL_DIR, EXPECTED_HEADERS
from sql_historial import listar_cola_revision_sql, guardar_en_cola_revision_sql, eliminar_de_cola_revision_sql, COLUMNAS

IDX_ARCHIVO = COLUMNAS.index("Archivo")

# Las 2 filas fantasma detectadas por el diagnóstico (en ColaRevision,
# cola="pendiente", sin PDF en corregir_manualmente/).
FANTASMA_PENDIENTE = [
    "F00097__EMAIL__marqueros@migasa.com__20260721_074122__ENDMAIL__.pdf",
    "F00102__EMAIL__marqueros@migasa.com__20260721_074125__ENDMAIL__.pdf",
]


def _archivos_en_disco(carpeta):
    if not os.path.exists(carpeta):
        return set()
    return {f for f in os.listdir(carpeta) if f.lower().endswith(".pdf")}


def _archivos_en_cola(cola):
    filas = listar_cola_revision_sql(cola)
    return {fila[IDX_ARCHIVO] for fila in filas}


def _cargar_historial_auto():
    """Archivo -> fila (en orden EXPECTED_HEADERS), última aparición gana."""
    path = os.path.join(HISTORIAL_DIR, "historial_facturas_auto.xlsx")
    if not os.path.exists(path):
        print(f"AVISO: no existe {path}, no se puede recuperar nada del historial.")
        return {}

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    n = len(EXPECTED_HEADERS)

    por_archivo = {}
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        archivo = row[0]
        if not isinstance(archivo, str) or not archivo.lower().endswith(".pdf"):
            continue
        fila = [("" if v is None else v) for v in (list(row) + [""] * n)[:n]]
        por_archivo[archivo] = fila

    wb.close()
    return por_archivo


def reparar_huerfanos(nombre, carpeta, cola, historial):
    disco = _archivos_en_disco(carpeta)
    en_bd = _archivos_en_cola(cola)
    huerfanos = sorted(disco - en_bd)

    print(f"\n=== {nombre} ({len(huerfanos)} huérfanos a reparar) ===")

    recuperados = fallback = fallidos = 0
    for archivo in huerfanos:
        fila = historial.get(archivo)
        origen = "historial"
        if fila is None:
            fila = [archivo] + ["-"] * (len(EXPECTED_HEADERS) - 1)
            origen = "placeholder"

        ok = guardar_en_cola_revision_sql(fila, cola)
        if not ok:
            fallidos += 1
            print(f"  [FALLO] {archivo}")
            continue

        if origen == "historial":
            recuperados += 1
        else:
            fallback += 1

    print(f"  Recuperados con datos del historial: {recuperados}")
    print(f"  Insertados como placeholder (sin datos en historial): {fallback}")
    print(f"  Fallidos: {fallidos}")


def limpiar_fantasmas():
    print(f"\n=== Filas fantasma en ColaRevision (sin PDF) ===")
    disco = _archivos_en_disco(CORREGIR_DIR)
    for archivo in FANTASMA_PENDIENTE:
        if archivo in disco:
            print(f"  [OMITIDO] {archivo} sí tiene PDF en disco, no se toca")
            continue
        ok = eliminar_de_cola_revision_sql(archivo)
        print(f"  {'[OK]' if ok else '[FALLO]'} eliminada {archivo}")


if __name__ == "__main__":
    historial = _cargar_historial_auto()
    print(f"Historial cargado: {len(historial)} archivos con datos recuperables.")

    reparar_huerfanos("Corregir manualmente", CORREGIR_DIR, "pendiente", historial)
    reparar_huerfanos("Incidencias", INCIDENCIAS_DIR, "incidencia", historial)
    limpiar_fantasmas()
