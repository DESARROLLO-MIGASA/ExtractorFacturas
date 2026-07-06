"""
Migración puntual: vuelca a la tabla SQL Server FacturasExaminadas las
facturas "examinada" ya acumuladas en el historial Excel (extracción
automática + correcciones manuales), antes de que la aplicación
empezara a guardarlas en SQL Server directamente.

Es una operación de un solo uso, pero segura de re-ejecutar: cada
factura se guarda por upsert sobre "Archivo" (ver sql_historial.py),
así que repetir la migración no crea duplicados.

Requisitos antes de ejecutar en serio (sin --dry-run):
- SQL_SERVER y SQL_DATABASE configurados en .env
- El usuario de Windows actual con permiso (al menos db_owner) sobre
  esa base de datos.

Uso:
    python migrar_historial_excel.py --dry-run   (solo cuenta y lista, no escribe nada)
    python migrar_historial_excel.py             (migra de verdad a SQL Server)
"""

import os
import sys
import argparse

from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "app_auto"))

from logic import (  # noqa: E402
    EXPECTED_HEADERS,
    CAMPOS_OBLIGATORIOS,
    CAMPOS_EXCLUIR_IMAGEN,
    limpiar_fila,
    FACTURAS_DIR,
)
from sql_historial import guardar_factura_examinada_sql  # noqa: E402

HISTORIAL_DIR = os.path.join(FACTURAS_DIR, "historial")
PATH_AUTO = os.path.join(HISTORIAL_DIR, "historial_facturas_auto.xlsx")
PATH_CORREGIDAS = os.path.join(HISTORIAL_DIR, "facturas_corregidas.xlsx")


def _fila_desde_dict(datos):
    fila = [
        "-" if datos.get(campo) is None else str(datos.get(campo))
        for campo in EXPECTED_HEADERS
    ]
    return limpiar_fila(fila)


def leer_bloques(path):
    """
    Lee un Excel con el formato de guardar_historial(): bloques repetidos
    de [título de lote] [cabecera] [una o varias filas de datos].
    La cabecera puede variar de un bloque a otro (el esquema de campos
    ha evolucionado con el tiempo), así que se relee en cada bloque.
    """
    if not os.path.exists(path):
        print(f"(no existe, se omite: {path})")
        return []

    wb = load_workbook(path, read_only=True)
    ws = wb.active

    filas = []
    cabecera_actual = None

    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue

        primera = str(row[0]).strip()

        if primera.startswith("Extracci"):  # "Extracción facturas | Fecha: ..."
            cabecera_actual = None
            continue

        if primera == "Archivo":
            cabecera_actual = [str(c).strip() if c is not None else "" for c in row]
            continue

        if cabecera_actual is None:
            continue

        datos = dict(zip(cabecera_actual, row))
        filas.append(_fila_desde_dict(datos))

    wb.close()
    return filas


def leer_tabla_plana(path):
    """Lee un Excel de tabla simple (una cabecera, filas de datos)."""
    if not os.path.exists(path):
        print(f"(no existe, se omite: {path})")
        return []

    wb = load_workbook(path, read_only=True)
    ws = wb.active

    filas_iter = ws.iter_rows(values_only=True)
    cabecera = [str(c).strip() if c is not None else "" for c in next(filas_iter)]

    filas = []
    for row in filas_iter:
        if not row or row[0] is None:
            continue
        datos = dict(zip(cabecera, row))
        filas.append(_fila_desde_dict(datos))

    wb.close()
    return filas


def clasificar_por_campos(fila):
    """
    Misma lógica que clasificar_factura() en app_auto/logic.py, pero sin
    la comprobación de facturas_corregidas.xlsx (eso se trata aparte,
    dándole prioridad como fuente definitiva de "examinada").
    """
    datos = dict(zip(EXPECTED_HEADERS, fila))

    todos_vacios = all(
        str(datos.get(campo, "-")).strip() == "-"
        for campo in EXPECTED_HEADERS
        if campo not in CAMPOS_EXCLUIR_IMAGEN
    )
    if todos_vacios:
        return "imagen"

    otros_obligatorios = [c for c in CAMPOS_OBLIGATORIOS if c != "PedidoCliente"]
    for campo in otros_obligatorios:
        if str(datos.get(campo, "-")).strip() == "-":
            return "manual"

    if str(datos.get("PedidoCliente", "-")).strip() == "-":
        return "reenviar_pedido"

    return "examinada"


def recopilar_examinadas():
    """
    Devuelve {archivo: (fila, origen)}, una sola fila por archivo.
    Las correcciones manuales (facturas_corregidas.xlsx) tienen
    prioridad sobre la extracción automática, por ser la versión
    definitiva revisada por una persona.
    """
    resultado = {}

    print(f"Leyendo extracción automática: {PATH_AUTO}")
    for fila in leer_bloques(PATH_AUTO):
        if clasificar_por_campos(fila) == "examinada":
            resultado[fila[0]] = (fila, "auto")

    print(f"Leyendo correcciones manuales: {PATH_CORREGIDAS}")
    for fila in leer_tabla_plana(PATH_CORREGIDAS):
        resultado[fila[0]] = (fila, "manual")

    return resultado


def main():
    parser = argparse.ArgumentParser(
        description="Migra facturas 'examinada' del historial Excel a SQL Server."
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Solo cuenta y lista, no escribe nada en SQL Server.",
    )
    args = parser.parse_args()

    print(f"FACTURAS_DIR: {FACTURAS_DIR}")

    examinadas = recopilar_examinadas()
    print(f"\nFacturas 'examinada' encontradas: {len(examinadas)}")

    if args.dry_run:
        for archivo, (_, origen) in list(examinadas.items())[:15]:
            print(f"  [{origen}] {archivo}")
        if len(examinadas) > 15:
            print(f"  ... y {len(examinadas) - 15} más")
        print("\nDry-run: no se ha escrito nada en SQL Server.")
        return

    ok, fallidas = 0, 0
    for i, (archivo, (fila, origen)) in enumerate(examinadas.items(), start=1):
        if guardar_factura_examinada_sql(fila, origen):
            ok += 1
        else:
            fallidas += 1
        if i % 50 == 0:
            print(f"  {i}/{len(examinadas)} procesadas...")

    print(f"\nMigración terminada: {ok} guardadas, {fallidas} fallidas.")
    if fallidas:
        print("Revisa los avisos anteriores: probablemente falte configurar SQL_SERVER/SQL_DATABASE en .env o permisos.")


if __name__ == "__main__":
    main()
