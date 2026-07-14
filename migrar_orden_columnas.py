"""
Migración puntual: reordena las columnas de los Excel de historial ya
guardados en disco para que coincidan con el nuevo orden de EXPECTED_HEADERS
en logic.py (antes alfabético, ahora agrupado por documento/comprador-
proveedor/referencia/importes).

Estos ficheros se leen por POSICIÓN (no por nombre de columna) en el resto
de la aplicación, así que un simple cambio de EXPECTED_HEADERS en el código
no basta: hay que reordenar también los datos ya escritos, o se leerían con
las columnas desplazadas.

Ficheros que toca (todos en FACTURAS_DIR/historial):
- historial_facturas_auto.xlsx / _imagenes.xlsx / _manual.xlsx
  (formato "bloques": título de lote + cabecera + filas, repetido)
- facturas_corregidas.xlsx / facturas_100_definitivas.xlsx
  (cabecera única + filas, con columnas extra al final: usuario/fecha/origen)

Antes de escribir nada hace una copia de cada fichero en
historial/_backup_orden_columnas/, y no la vuelve a pisar si ya existe (para
poder repetir la migración sin perder el respaldo original).

Uso:
    python migrar_orden_columnas.py --dry-run   (solo informa, no escribe nada)
    python migrar_orden_columnas.py             (migra de verdad)
"""

import os
import shutil
import argparse

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from logic import HISTORIAL_DIR, _guardar_excel_atomico

# Orden ANTERIOR (alfabético) tal cual está escrito en los ficheros ya
# guardados. Se deja fijo aquí, independiente de logic.py, para que esta
# migración siga siendo correcta aunque en el futuro EXPECTED_HEADERS
# cambie otra vez.
OLD_HEADERS = [
    "Archivo", "BaseImp", "BaseIRPF", "Buyer", "Empresa", "FEscaneo",
    "FFactura", "FOperacion", "ImporIVA", "Moneda", "NombreProveedor",
    "NumeroFactura", "PedidoCliente", "Proveedor", "TipoIVA", "TipoIVA2",
    "TipoIVA3", "TotalFact",
]

# Orden NUEVO acordado: identificación del documento -> comprador/proveedor
# (CIF junto a su nombre) -> referencia -> importes.
NEW_HEADERS = [
    "Archivo", "NumeroFactura", "Buyer", "Empresa", "Proveedor",
    "NombreProveedor", "PedidoCliente", "BaseImp", "BaseIRPF", "TipoIVA",
    "TipoIVA2", "TipoIVA3", "ImporIVA", "TotalFact", "Moneda", "FFactura",
    "FOperacion", "FEscaneo",
]

assert sorted(OLD_HEADERS) == sorted(NEW_HEADERS), "OLD_HEADERS y NEW_HEADERS deben tener los mismos campos"

BACKUP_DIR = os.path.join(HISTORIAL_DIR, "_backup_orden_columnas")

FICHEROS_BLOQUES = [
    "historial_facturas_auto.xlsx",
    "historial_facturas_imagenes.xlsx",
    "historial_facturas_manual.xlsx",
]

FICHEROS_PLANOS = [
    "facturas_corregidas.xlsx",
    "facturas_100_definitivas.xlsx",
]

FILL_VACIO = PatternFill("solid", fgColor="E7E6E6")
FONT_VACIO = Font(color="666666")
FONT_NORMAL = Font(color="000000")


def _backup(path):
    os.makedirs(BACKUP_DIR, exist_ok=True)
    destino = os.path.join(BACKUP_DIR, os.path.basename(path))
    if os.path.exists(destino):
        return destino, False
    shutil.copy2(path, destino)
    return destino, True


def _nuevo_orden_completo(cabecera_actual):
    """
    A partir de la cabecera real de un bloque/fichero (puede traer columnas
    extra al final, p.ej. "Origen"/"UsuarioDefinitiva"/"FechaDefinitiva"),
    calcula el nuevo orden completo: los 18 campos núcleo en NEW_HEADERS,
    seguidos de cualquier columna extra en el mismo orden relativo en que
    ya estaba.
    """
    extra = [h for h in cabecera_actual if h not in OLD_HEADERS]
    return NEW_HEADERS + extra


def _migrar_hoja_bloques(ws, stats):
    cabecera_actual = None
    permutacion = None  # permutacion[i] = índice ANTIGUO que va en la posición NUEVA i

    for row in ws.iter_rows():
        primera = row[0].value

        if primera is None:
            continue

        primera_str = str(primera).strip()

        if primera_str.startswith("Extracci"):
            cabecera_actual = None
            permutacion = None
            continue

        if primera_str == "Archivo":
            cabecera_actual = [str(c.value).strip() if c.value is not None else "" for c in row]
            nuevo_orden = _nuevo_orden_completo(cabecera_actual)
            permutacion = [cabecera_actual.index(h) for h in nuevo_orden]

            for i, celda in enumerate(row[:len(nuevo_orden)]):
                celda.value = nuevo_orden[i]

            stats["bloques"] += 1
            continue

        if permutacion is None:
            continue

        valores_originales = [c.value for c in row[:len(permutacion)]]
        nuevos_valores = [valores_originales[permutacion[i]] for i in range(len(permutacion))]

        for i, celda in enumerate(row[:len(nuevos_valores)]):
            celda.value = nuevos_valores[i]
            if str(nuevos_valores[i] if nuevos_valores[i] is not None else "").strip() == "-":
                celda.fill = FILL_VACIO
                celda.font = FONT_VACIO
            else:
                celda.font = FONT_NORMAL

        stats["filas"] += 1


def _migrar_hoja_plana(ws, stats):
    filas = list(ws.iter_rows())
    if not filas:
        return

    fila_cabecera = filas[0]
    cabecera_actual = [str(c.value).strip() if c.value is not None else "" for c in fila_cabecera]
    nuevo_orden = _nuevo_orden_completo(cabecera_actual)
    permutacion = [cabecera_actual.index(h) for h in nuevo_orden]

    for i, celda in enumerate(fila_cabecera[:len(nuevo_orden)]):
        celda.value = nuevo_orden[i]
    stats["bloques"] += 1

    for row in filas[1:]:
        if row[0].value is None:
            continue

        valores_originales = [c.value for c in row[:len(permutacion)]]
        nuevos_valores = [valores_originales[permutacion[i]] for i in range(len(permutacion))]

        for i, celda in enumerate(row[:len(nuevos_valores)]):
            celda.value = nuevos_valores[i]
            if str(nuevos_valores[i] if nuevos_valores[i] is not None else "").strip() == "-":
                celda.fill = FILL_VACIO
                celda.font = FONT_VACIO
            else:
                celda.font = FONT_NORMAL

        stats["filas"] += 1


def migrar_fichero(nombre, es_bloques, dry_run):
    path = os.path.join(HISTORIAL_DIR, nombre)

    if not os.path.exists(path):
        print(f"{nombre}: no existe, se omite.")
        return

    stats = {"bloques": 0, "filas": 0}

    if dry_run:
        wb = load_workbook(path, read_only=True)
        ws = wb.active
        cabeceras = 0
        filas = 0
        for row in ws.iter_rows(values_only=True):
            if not row or row[0] is None:
                continue
            if es_bloques and str(row[0]).strip().startswith("Extracci"):
                continue
            if row[0] == "Archivo":
                cabeceras += 1
            else:
                filas += 1
        wb.close()
        print(f"{nombre}: [dry-run] {cabeceras} cabecera(s), {filas} fila(s) de datos a reordenar.")
        return

    _, copiado = _backup(path)
    print(f"{nombre}: copia de seguridad {'creada' if copiado else 'ya existía'} en {BACKUP_DIR}")

    wb = load_workbook(path)
    ws = wb.active

    if es_bloques:
        _migrar_hoja_bloques(ws, stats)
    else:
        _migrar_hoja_plana(ws, stats)

    _guardar_excel_atomico(wb, path)
    print(f"{nombre}: {stats['bloques']} cabecera(s) y {stats['filas']} fila(s) reordenadas.")


def main():
    parser = argparse.ArgumentParser(
        description="Reordena las columnas de los Excel de historial al nuevo orden de EXPECTED_HEADERS."
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Solo informa de lo que se migraría, no escribe nada.",
    )
    args = parser.parse_args()

    print(f"HISTORIAL_DIR: {HISTORIAL_DIR}")
    print(f"Orden antiguo: {OLD_HEADERS}")
    print(f"Orden nuevo:   {NEW_HEADERS}\n")

    for nombre in FICHEROS_BLOQUES:
        migrar_fichero(nombre, es_bloques=True, dry_run=args.dry_run)

    for nombre in FICHEROS_PLANOS:
        migrar_fichero(nombre, es_bloques=False, dry_run=args.dry_run)

    if args.dry_run:
        print("\nDry-run: no se ha escrito nada.")
    else:
        print("\nMigración terminada.")


if __name__ == "__main__":
    main()
