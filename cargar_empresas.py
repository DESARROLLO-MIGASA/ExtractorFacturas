"""
Carga puntual: vuelca a la tabla SQL EmpresasClasificadas las empresas de
"Empresas granel.xlsx" (Granel) y "Empresas envasado.xlsb" (Envasado), con
su CIF, código de empresa, nombre y si están activas o no.

Un mismo CIF puede aparecer en los dos ficheros (empresa que opera a granel
y envasado); en ese caso se guardan dos filas, una por clasificación. Es
seguro re-ejecutar: cada fila se guarda por upsert sobre (CIF,
Clasificacion), así que repetir la carga no crea duplicados.

Usa el mismo motor configurado en .env para el resto de la app
(SQL_ENGINE=sqlserver o sqlite, ver sql_historial.py).

Uso:
    python cargar_empresas.py --dry-run   (solo cuenta y lista, no escribe nada)
    python cargar_empresas.py             (carga de verdad)
"""

import os
import argparse

import openpyxl
from pyxlsb import open_workbook as open_workbook_xlsb

from sql_historial import _conectar, MOTOR, FECHA_ACTUAL_SQL

_ROOT = os.path.dirname(os.path.abspath(__file__))

FICHEROS = [
    (os.path.join(_ROOT, "Empresas granel.xlsx"), "Granel"),
    (os.path.join(_ROOT, "Empresas envasado.xlsb"), "Envasado"),
]

TABLA = "EmpresasClasificadas"


def _filas_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    return list(ws.iter_rows(values_only=True))


def _filas_xlsb(path):
    with open_workbook_xlsb(path) as wb:
        with wb.get_sheet(1) as sheet:
            return [tuple(c.v for c in fila) for fila in sheet.rows()]


def leer_empresas(path, clasificacion):
    """
    Lee un fichero exportado de Business Central (vista/edición de la tabla
    Empresa). La fila 0 es el título de la vista ("Vista - Empresa" /
    "Editar - Empresa"), la fila 1 trae la cabecera real y desde la fila 2
    vienen los datos. El orden de columnas difiere entre el .xlsx de Granel
    y el .xlsb de Envasado, así que se busca por nombre de columna.
    """
    ext = os.path.splitext(path)[1].lower()
    filas = _filas_xlsb(path) if ext == ".xlsb" else _filas_xlsx(path)

    headers = filas[1]
    idx = {h: i for i, h in enumerate(headers)}

    empresas = []
    for fila in filas[2:]:
        cif = str(fila[idx["CIF/NIF"]] or "").strip()
        if not cif:
            continue
        codigo = str(fila[idx["Código Empresa"]] or "").strip()
        nombre = str(fila[idx["Nombre"]] or "").strip()
        activa = str(fila[idx["Empresa Activa"]] or "").strip().lower() in ("sí", "si")
        empresas.append((clasificacion, cif, codigo, nombre, activa))

    return empresas


def _crear_tabla_si_no_existe(cursor):
    if MOTOR == "sqlite":
        cursor.execute(f"""
            CREATE TABLE IF NOT EXISTS {TABLA} (
                Id INTEGER PRIMARY KEY AUTOINCREMENT,
                Clasificacion TEXT NOT NULL,
                CIF TEXT NOT NULL,
                CodigoEmpresa TEXT NOT NULL,
                NombreEmpresa TEXT NOT NULL,
                Activa INTEGER NOT NULL DEFAULT 0,
                FechaCarga TEXT NOT NULL DEFAULT ({FECHA_ACTUAL_SQL}),
                UNIQUE(CIF, Clasificacion)
            )
        """)
        return

    cursor.execute(f"""
        IF NOT EXISTS (SELECT 1 FROM sys.tables WHERE name = '{TABLA}')
        CREATE TABLE {TABLA} (
            Id INT IDENTITY(1,1) PRIMARY KEY,
            Clasificacion NVARCHAR(20) NOT NULL,
            CIF NVARCHAR(30) NOT NULL,
            CodigoEmpresa NVARCHAR(20) NOT NULL,
            NombreEmpresa NVARCHAR(200) NOT NULL,
            Activa BIT NOT NULL DEFAULT 0,
            FechaCarga DATETIME NOT NULL DEFAULT GETDATE(),
            CONSTRAINT UQ_{TABLA}_CIF_Clasificacion UNIQUE (CIF, Clasificacion)
        )
    """)


def guardar_empresa(cursor, clasificacion, cif, codigo, nombre, activa):
    activa_val = 1 if activa else 0

    cursor.execute(
        f"UPDATE {TABLA} SET CodigoEmpresa = ?, NombreEmpresa = ?, Activa = ?, FechaCarga = {FECHA_ACTUAL_SQL} "
        f"WHERE CIF = ? AND Clasificacion = ?",
        (codigo, nombre, activa_val, cif, clasificacion),
    )
    if cursor.rowcount == 0:
        cursor.execute(
            f"INSERT INTO {TABLA} (Clasificacion, CIF, CodigoEmpresa, NombreEmpresa, Activa) "
            f"VALUES (?, ?, ?, ?, ?)",
            (clasificacion, cif, codigo, nombre, activa_val),
        )


def main():
    parser = argparse.ArgumentParser(
        description='Carga "Empresas granel.xlsx" y "Empresas envasado.xlsb" en la tabla EmpresasClasificadas.'
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Solo cuenta y lista, no escribe nada en la base de datos.",
    )
    args = parser.parse_args()

    todas = []
    for path, clasificacion in FICHEROS:
        print(f"Leyendo {clasificacion}: {path}")
        empresas = leer_empresas(path, clasificacion)
        print(f"  {len(empresas)} empresas")
        todas.extend(empresas)

    activas = sum(1 for _, _, _, _, a in todas if a)
    print(f"\nTotal filas a guardar: {len(todas)} ({activas} activas, {len(todas) - activas} inactivas)")

    if args.dry_run:
        for clasificacion, cif, codigo, nombre, activa in todas[:15]:
            marca = "" if activa else " [INACTIVA]"
            print(f"  [{clasificacion}] {cif} - ({codigo}) {nombre}{marca}")
        if len(todas) > 15:
            print(f"  ... y {len(todas) - 15} más")
        print(f"\nDry-run: no se ha escrito nada en {MOTOR}.")
        return

    with _conectar() as conn:
        cursor = conn.cursor()
        _crear_tabla_si_no_existe(cursor)
        conn.commit()

        for i, (clasificacion, cif, codigo, nombre, activa) in enumerate(todas, start=1):
            guardar_empresa(cursor, clasificacion, cif, codigo, nombre, activa)
            if i % 100 == 0:
                print(f"  {i}/{len(todas)} procesados...")

        conn.commit()

    print(f"\nCarga terminada: {len(todas)} empresas guardadas en {TABLA} ({MOTOR}).")


if __name__ == "__main__":
    main()
